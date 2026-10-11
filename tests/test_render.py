"""Renderer tests: HTML, PDF and DOCX output must be real and editable."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from painter.render import HtmlRenderOptions, render_docx, render_document, render_pdf
from painter.render.pdf_renderer import extract_pdf_text, pdf_page_count
from painter.schemas import (
    Alignment,
    BBox,
    DocumentIR,
    DocumentMetadata,
    DocumentSource,
    Element,
    ElementType,
    ImageContent,
    PageFormat,
    PageIR,
    TableCell,
    TableContent,
    Typography,
)


def build_document(tmp_path: Path) -> DocumentIR:
    import cv2
    import numpy as np

    asset = tmp_path / "logo.png"
    cv2.imwrite(str(asset), np.full((60, 120, 3), 120, dtype=np.uint8))

    heading = Element(
        id="el_heading",
        type=ElementType.HEADING,
        bbox=BBox(x1=100, y1=100, x2=900, y2=160),
        confidence=0.95,
        text="QUARTERLY REPORT & SUMMARY",
        alignment=Alignment.LEFT,
        typography=Typography(font_size_pt=20, font_weight=700, estimated=False),
        order_index=0,
    )
    paragraph = Element(
        id="el_paragraph",
        type=ElementType.PARAGRAPH,
        bbox=BBox(x1=100, y1=200, x2=900, y2=320),
        confidence=0.9,
        text="Revenue grew across all segments.\nCosts remained stable.",
        typography=Typography(font_size_pt=11, line_height_pt=14, estimated=False),
        order_index=1,
    )
    logo = Element(
        id="el_logo",
        type=ElementType.LOGO,
        bbox=BBox(x1=100, y1=360, x2=220, y2=420),
        confidence=0.8,
        image=ImageContent(asset_path=str(asset), asset_origin="extracted"),
        order_index=2,
    )
    table = Element(
        id="el_table",
        type=ElementType.TABLE,
        bbox=BBox(x1=100, y1=460, x2=900, y2=660),
        confidence=0.85,
        text="Item | Amount",
        table=TableContent(
            rows=2,
            cols=2,
            cells=[
                TableCell(row=0, col=0, text="Item", is_header=True),
                TableCell(row=0, col=1, text="Amount", is_header=True),
                TableCell(row=1, col=0, text="Consulting"),
                TableCell(row=1, col=1, text="1250000"),
            ],
            column_widths_mm=[100.0, 60.0],
        ),
        order_index=3,
    )
    line = Element(
        id="el_line",
        type=ElementType.LINE,
        bbox=BBox(x1=100, y1=700, x2=900, y2=702),
        confidence=0.7,
        order_index=4,
    )
    page = PageIR(
        index=0,
        width_px=1000,
        height_px=1400,
        page_format=PageFormat(name="custom", width_mm=210, height_mm=297),
        elements=[heading, paragraph, logo, table, line],
        reading_order=[e.id for e in (heading, paragraph, logo, table, line)],
        image_path="pages/page_00.png",
    )
    document = DocumentIR(
        document_id="doc_render",
        source=DocumentSource(
            filename="report.png",
            stored_path=str(tmp_path / "report.png"),
            sha256="deadbeef",
            page_count=1,
        ),
        metadata=DocumentMetadata(title="Test"),
        pages=[page],
    )
    document.recompute_stats()
    return document


def test_html_contains_positioned_elements(tmp_path):
    document = build_document(tmp_path)
    markup = render_document(document, HtmlRenderOptions(include_element_ids=True))
    assert "<!DOCTYPE html>" in markup
    assert "@page { size: 210.00mm 297.00mm" in markup
    assert 'data-element-id="el_heading"' in markup
    assert "position:absolute" in markup
    assert "QUARTERLY REPORT &amp; SUMMARY" in markup  # escaped
    assert "border-collapse:collapse" in markup  # native table markup
    assert "el-line" in markup


def test_html_asset_url_prefix(tmp_path):
    document = build_document(tmp_path)
    markup = render_document(document, HtmlRenderOptions(asset_url_prefix="/api/documents/x/files"))
    assert "/api/documents/x/files/" in markup


def test_html_reports_missing_asset(tmp_path):
    document = build_document(tmp_path)
    document.pages[0].elements[2].image.asset_path = None
    markup = render_document(document)
    assert "no asset" in markup


def test_html_global_offset_and_font_scale(tmp_path):
    document = build_document(tmp_path)
    base = render_document(document)
    shifted = render_document(document, HtmlRenderOptions(global_offset_mm=(5.0, 7.0), font_scale=1.2))
    assert base != shifted
    assert "left:26" in shifted or "left:27" in shifted  # 100px * 0.21mm + 5mm
    assert "font-size: 13.2pt" in shifted  # 11pt * 1.2


def test_pdf_export_has_text(tmp_path):
    document = build_document(tmp_path)
    result = render_pdf(document, tmp_path / "out.pdf")
    assert result.path.is_file()
    assert result.engine in {"weasyprint", "chromium"}
    assert result.has_text, "the PDF must contain extractable text, not a page image"
    assert "QUARTERLY" in extract_pdf_text(result.path).upper()
    assert result.pages == 1


def test_pdf_page_count_and_orientation(tmp_path):
    document = build_document(tmp_path)
    document.pages[0].page_format.width_mm = 297
    document.pages[0].page_format.height_mm = 210
    result = render_pdf(document, tmp_path / "landscape.pdf")
    assert pdf_page_count(result.path) == 1


def test_docx_export_is_editable(tmp_path):
    document = build_document(tmp_path)
    result = render_docx(document, tmp_path / "out.docx", asset_base=tmp_path)
    assert result.path.is_file()
    assert result.paragraphs >= 2
    assert result.tables == 1
    assert result.pictures >= 1
    assert result.editable_characters > 30

    import docx

    reopened = docx.Document(str(result.path))
    texts = [paragraph.text for paragraph in reopened.paragraphs if paragraph.text.strip()]
    assert any("QUARTERLY" in text for text in texts)
    assert any("Revenue grew" in text for text in texts)
    assert len(reopened.tables) == 1
    assert reopened.tables[0].cell(0, 0).text == "Item"
    assert reopened.tables[0].cell(1, 1).text == "1250000"


def test_docx_does_not_embed_a_full_page_image(tmp_path):
    document = build_document(tmp_path)
    result = render_docx(document, tmp_path / "check.docx", asset_base=tmp_path)
    with zipfile.ZipFile(result.path) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    # only the logo and the drawn line — never a rasterised page
    assert len(media) <= 4, media


def test_docx_anchors_pictures(tmp_path):
    document = build_document(tmp_path)
    result = render_docx(document, tmp_path / "anchored.docx", asset_base=tmp_path)
    assert result.anchored_pictures >= 1
    with zipfile.ZipFile(result.path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "<wp:anchor" in document_xml
    assert "positionH" in document_xml


def test_docx_flags_shape_fallback(tmp_path):
    document = build_document(tmp_path)
    result = render_docx(document, tmp_path / "shapes.docx", asset_base=tmp_path)
    assert result.shape_fallbacks >= 1
    assert result.warnings


def test_docx_handles_missing_asset(tmp_path):
    document = build_document(tmp_path)
    document.pages[0].elements[2].image.asset_path = str(tmp_path / "ghost.png")
    result = render_docx(document, tmp_path / "missing.docx", asset_base=tmp_path)
    assert any("missing asset" in warning for warning in result.warnings)


def test_render_empty_document_does_not_crash(tmp_path):
    document = DocumentIR(document_id="empty")
    markup = render_document(document)
    assert "DOCTYPE" in markup
    result = render_pdf(document, tmp_path / "empty.pdf")
    assert result.path.is_file()
